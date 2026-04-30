import base64, io, re, os, tempfile, json
from pathlib import Path
from typing import Optional, List, Dict, Any, Union
from fastapi import APIRouter, Form, UploadFile, File, HTTPException
from openai import AsyncOpenAI
import httpx, openpyxl, PyPDF2
from docx import Document
from config import load_settings, save_settings, logger

router = APIRouter(prefix="/api/ai", tags=["Генерация ТК"])

CSV_HEADERS = [
    "Класс", "Подкласс", "Нормализованный код модели",
    "Элемент", "Подэлемент", "Наименование операции",
    "Краткое содержание работ", "Вид ТОиР", "Периодичность",
    "Норма времени, часов", "Количество исполнителей",
    "Профессия/Квалификация", "Трудоёмкость, человеко/часов",
    "Наименование ТМЦ", "Количество ТМЦ", "Единицы измерения ТМЦ",
    "Наименование инструменты", "Средства индивидуальной защиты",
    "Требования по безопасности"
]

# ✅ Только рабочие модели по умолчанию
DEFAULT_WORKING_MODELS = [
    "openai/gpt-4o-2024-08-06",
    "openai/gpt-4o-mini",
    "meta-llama/llama-3.3-70b-instruct",
    "mistralai/mistral-large-2411",
    "openrouter/free",
    "z-ai/glm-4.5-air:free",
    "inclusionai/ling-2.6-1t:free",
    "openai/gpt-oss-120b:free"
]


def extract_text_from_file(file_path: str) -> str:
    ext = Path(file_path).suffix.lower()
    text = ""
    try:
        if ext == ".pdf":
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text: text += page_text + "\n"
        elif ext == ".docx":
            doc = Document(file_path)
            text = "\n".join([p.text for p in doc.paragraphs])
        elif ext in (".txt", ".csv"):
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        else:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
    except Exception as e:
        logger.error(f"Ошибка чтения файла {file_path}: {e}")
        text = f"[Ошибка: {e}]"
    return text


def parse_ai_table_response(text_response: str) -> List[List[str]]:
    if not text_response: return []
    rows = []
    for line in text_response.strip().split("\n"):
        line = line.strip()
        if not line or line.startswith("#") or "Элемент|Подэлемент" in line: continue
        parts = [p.strip() for p in line.split("|")]
        if len(parts) >= 2:
            while len(parts) < len(CSV_HEADERS): parts.append("")
            rows.append(parts[:len(CSV_HEADERS)])
    return rows


def create_xlsx(headers: List[str], rows: List[List[str]],
                class_val: str = "", subclass_val: str = "", model_code: str = "") -> bytes:
    wb = openpyxl.Workbook();
    ws = wb.active;
    ws.title = "Технологическая карта"
    header_font = openpyxl.styles.Font(name="Arial", bold=True, size=10, color="FFFFFF")
    header_fill = openpyxl.styles.PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_alignment = openpyxl.styles.Alignment(vertical="top", wrap_text=True)
    thin_border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style="thin"), right=openpyxl.styles.Side(style="thin"),
        top=openpyxl.styles.Side(style="thin"), bottom=openpyxl.styles.Side(style="thin")
    )
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font;
        cell.fill = header_fill;
        cell.alignment = header_alignment;
        cell.border = thin_border
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = cell_alignment;
            cell.border = thin_border;
            cell.font = openpyxl.styles.Font(name="Arial", size=9)
            if col_idx == 1 and not value and class_val:
                cell.value = class_val
            elif col_idx == 2 and not value and subclass_val:
                cell.value = subclass_val
            elif col_idx == 3 and not value and model_code:
                cell.value = model_code
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 18
    ws.freeze_panes = "A2";
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO();
    wb.save(buf);
    buf.seek(0)
    return buf.read()


async def call_ai(messages: List[Dict[str, str]], settings: Dict[str, Any]) -> str:
    provider = settings.get("provider", "openrouter")
    api_key = settings.get("api_key", "").strip()
    model = settings.get("model", "openai/gpt-4o-2024-08-06")  # ✅ Рабочая дефолтная модель

    if not api_key:
        raise HTTPException(status_code=400, detail="API ключ не установлен")

    if provider == "openai":
        client = AsyncOpenAI(api_key=api_key)
        resp = await client.chat.completions.create(model=model or "gpt-4o", messages=messages, temperature=0.3,
                                                    max_tokens=settings.get("max_tokens", 3000))
        return resp.choices[0].message.content
    elif provider == "openrouter":
        async with httpx.AsyncClient(timeout=120.0) as client:
            resp = await client.post("https://openrouter.ai/api/v1/chat/completions",
                                     headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json",
                                              "HTTP-Referer": "http://localhost:8000",
                                              "X-Title": "Unified Engineering Portal"},
                                     json={"model": model, "messages": messages, "temperature": 0.3,
                                           "max_tokens": settings.get("max_tokens", 3000)})
            if resp.status_code != 200:
                logger.error(f"OpenRouter {resp.status_code}: {resp.text}")
                raise HTTPException(status_code=resp.status_code, detail=f"OpenRouter: {resp.text[:200]}")
            return resp.json()["choices"][0]["message"]["content"]
    elif provider == "yandex":
        async with httpx.AsyncClient(timeout=120.0) as client:
            resp = await client.post("https://llm.api.cloud.yandex.net/foundationModels/v1/openai/chat/completions",
                                     headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                                     json={"model": model or "yandexgpt/latest", "messages": messages,
                                           "temperature": 0.3, "max_tokens": settings.get("max_tokens", 3000)})
            if resp.status_code != 200: raise HTTPException(status_code=resp.status_code,
                                                            detail=f"YandexGPT: {resp.text[:200]}")
            return resp.json()["choices"][0]["message"]["content"]
    elif provider == "gigachat":
        async with httpx.AsyncClient(timeout=120.0) as client:
            resp = await client.post("https://gigachat.devices.sberbank.ru/api/v1/chat/completions",
                                     headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                                     json={"model": model or "GigaChat-Pro", "messages": messages, "temperature": 0.3,
                                           "max_tokens": settings.get("max_tokens", 3000)})
            if resp.status_code != 200: raise HTTPException(status_code=resp.status_code,
                                                            detail=f"GigaChat: {resp.text[:200]}")
            return resp.json()["choices"][0]["message"]["content"]
    else:
        raise HTTPException(status_code=400, detail=f"Неизвестный провайдер: {provider}")


@router.get("/settings")
async def get_settings(): return load_settings()


@router.post("/settings")
async def update_settings(settings: Dict[str, Any]):
    current = load_settings();
    current.update(settings)
    if save_settings(current): return {"status": "ok"}
    raise HTTPException(status_code=500, detail="Ошибка сохранения")


@router.post("/chat")
async def generate_tk(
        message: str = Form(...), model_name: str = Form(""), equipment_class: str = Form(""), subclass: str = Form(""),
        file: Union[UploadFile, str, None] = File(default=None),
        model: Optional[str] = Form(None, alias="model")  # ← Алиас для приёма "model" из формы
):
    try:
        settings = load_settings()

        # 🔑 Переопределение модели из запроса
        if model and model.strip():
            settings["model"] = model.strip()
            logger.info(f"🔄 Модель переопределена из запроса: {model}")

        ai_settings = settings.get("ai", {}) if "ai" in settings else settings
        api_key = ai_settings.get("api_key", "").strip()
        if not api_key:
            raise HTTPException(status_code=400, detail="API ключ не установлен")

        # Нормализация файла
        if isinstance(file, str) and not file.strip():
            file = None
        elif isinstance(file, UploadFile) and (not file.filename or not file.filename.strip()):
            file = None

        file_text = ""
        if file and isinstance(file, UploadFile):
            temp_path = os.path.join(tempfile.gettempdir(), os.path.basename(file.filename))
            with open(temp_path, "wb") as f:
                f.write(await file.read())
            file_text = extract_text_from_file(temp_path)
            try:
                os.remove(temp_path)
            except:
                pass

        file_instruction = f"""📄 ТЕХНИЧЕСКИЙ ПАСПОРТ для модели "{model_name}".
Используй ТОЛЬКО информацию из документа.
СОДЕРЖИМОЕ (первые 40000 символов):
{file_text[:40000]}
---""" if file_text.strip() else f"""📋 ПАСПОРТ НЕ ЗАГРУЖЕН. Используй общие знания по модели "{model_name}"."""

        base_prompt = ai_settings.get("master_prompt", "")
        system_content = base_prompt.format(
            file_instruction=file_instruction) if "{file_instruction}" in base_prompt else f"{file_instruction}\n\n{base_prompt}"
        system_content += """
ОТВЕТ В ФОРМАТЕ:
[ТЕКСТ_ОТВЕТ]Краткое описание[/ТЕКСТ_ОТВЕТ]
[ТАБЛИЦА]Элемент|Подэлемент|...|Требования по безопасности[/ТАБЛИЦА]"""

        messages = [{"role": "system", "content": system_content}, {"role": "user",
                                                                    "content": f"Модель: {model_name}\nКласс: {equipment_class}\nПодкласс: {subclass}\nЗапрос: {message}"}]

        ai_response = await call_ai(messages, ai_settings)

        # 🔑 ЗАЩИТА: проверка ai_response на None
        if not ai_response or not isinstance(ai_response, str):
            logger.error(f"Пустой или некорректный ответ от ИИ: {type(ai_response)}")
            raise HTTPException(status_code=500, detail="Пустой ответ от нейросети")

        # Парсинг ответа
        text_match = re.search(r"\[ТЕКСТ_ОТВЕТ\](.*?)\[/ТЕКСТ_ОТВЕТ\]", ai_response, re.DOTALL)
        table_match = re.search(r"\[ТАБЛИЦА\](.*?)\[/ТАБЛИЦА\]", ai_response, re.DOTALL)

        text_part = text_match.group(1).strip() if text_match else ai_response[:400]
        table_part = table_match.group(1).strip() if table_match else ""
        rows = parse_ai_table_response(table_part)

        xlsx_data = None
        if rows:
            table_headers = CSV_HEADERS[3:]
            xlsx_bytes = create_xlsx(table_headers, rows, equipment_class, subclass, model_name)
            xlsx_data = base64.b64encode(xlsx_bytes).decode()

        return {"text": text_part, "table_rows": rows, "xlsx_file": xlsx_data,
                "xlsx_filename": f"ТК_{model_name or 'модель'}.xlsx", "status": "success"}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Ошибка генерации ТК: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Ошибка: {str(e)[:200]}")


@router.get("/table_template")
async def get_table_template(): return {"headers": CSV_HEADERS}